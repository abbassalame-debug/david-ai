"""
restaurant_weights.py -- Job 1: reads the manager's weights workbook (`فوارق الشك مع الاوزان <date>.xlsx`), the PHYSICAL truth about the meat.

The manager's rule: the weights workbook wins over the company's consumption table, which the draft reads. The workbook has one tab per day (named
`dd-mm-yyyy`; some Sep tabs are stale, so the tab is found by its date) and holds, per meat: every spit's weight ("الوزن السيخ"), the opening weigh-in and
the end-of-shift closing weigh-in ("قبل الشك": chicken trim D / main E; beef trim O / tallow P / main Q). Read-only. Nothing is guessed: a figure that is not
there is None, and a sum that does not equal its parts is a flag.

    py restaurant_weights.py 2026-10-06            print what the workbook says for the day
"""
import datetime
import glob
import os
import re
import sys
from decimal import Decimal

import openpyxl

import menu_tools

FRIDGE_LARGE_KG = Decimal("5")
SPIT_LABEL = "الوزن السيخ"
BEEF_FIRST_COL = 14


def _dec(v):
    if isinstance(v, bool) or not isinstance(v, (int, float)):
        return None
    return Decimal(str(round(float(v), 3)))


def _sheet_date(name):
    m = re.match(r"\s*0?(\d+)-0?(\d+)-(\d{4})", name or "")
    return (int(m.group(3)), int(m.group(2)), int(m.group(1))) if m else None


def _name_date(path):
    m = re.search(r"(\d{1,2})-(\d{1,2})-(\d{4})(?!.*\d{1,2}-\d{1,2}-\d{4})", os.path.basename(path))
    return (int(m.group(3)), int(m.group(2)), int(m.group(1))) if m else (0, 0, 0)


def _tab(wb, day):
    y, m, d = map(int, day.split("-"))
    return next((wb[n] for n in wb.sheetnames if _sheet_date(n) == (y, m, d)), None)


def find_workbook(day, restaurant_root=None, extra_dirs=None):
    """(path, note): the latest-dated weights workbook that has a tab for the day, anywhere in the Demo Bistro month folders (and any extra folder, e.g. the
    day's upload inbox), or (None, note)."""
    root = restaurant_root or menu_tools.RESTAURANT_ROOT
    pool = [p for p in glob.glob(os.path.join(root, "*-20??", "*", "فوارق*.xlsx")) if not os.path.basename(p).startswith("~$")]
    for d in extra_dirs or []:
        pool += [p for p in glob.glob(os.path.join(d, "فوارق*.xlsx")) if not os.path.basename(p).startswith("~$")]
    pool = sorted(set(pool))
    for p in sorted(pool, key=lambda p: (_name_date(p), os.path.getmtime(p), p), reverse=True):
        try:
            wb = openpyxl.load_workbook(p, data_only=True, read_only=False)
        except Exception:
            continue
        if _tab(wb, day) is not None:
            return p, f"the latest-dated weights workbook that has a tab for {day} ({os.path.basename(p)})"
    return None, f"no weights workbook with a tab for {day} was found"


def _find(ws, text):
    return [c for row in ws.iter_rows() for c in row if isinstance(c.value, str) and text in c.value]


def _weigh_row(ws, r):
    return {"chicken": {"trim_kg": _dec(ws.cell(r, 4).value), "main_kg": _dec(ws.cell(r, 5).value)},
            "beef": {"trim_kg": _dec(ws.cell(r, 15).value), "tallow_kg": _dec(ws.cell(r, 16).value), "main_kg": _dec(ws.cell(r, 17).value)}}


def read_day(path, day):
    """What the workbook says for the day: {"sheet", "spits": {meat: [kg]}, "spit_sum_cell": {meat: kg|None}, "opening", "closing", "flags": [..]}."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = _tab(wb, day)
    if ws is None:
        raise ValueError(f"{os.path.basename(path)} has no tab for {day}")
    out = {"sheet": ws.title, "spits": {"chicken": [], "beef": []}, "spit_sum_cell": {}, "opening": None, "closing": None, "flags": []}
    for c in _find(ws, SPIT_LABEL):
        if c.column < BEEF_FIRST_COL:
            v = _dec(ws.cell(c.row, c.column + 2).value)
            if v is not None:
                out["spits"]["chicken"].append(v)
        else:
            v = _dec(ws.cell(c.row, c.column + 3).value)
            if v is not None:
                out["spits"]["beef"].append(v)
    for lab, meat in (("مجموع السياخ دجاج", "chicken"), ("مجموع السياخ لحم", "beef")):
        cs = _find(ws, lab)
        out["spit_sum_cell"][meat] = _dec(ws.cell(cs[0].row + 2, cs[0].column).value) if cs else None
        parts = sum(out["spits"][meat], Decimal(0))
        printed = out["spit_sum_cell"][meat]
        if printed is not None and out["spits"][meat] and abs(printed - parts) > Decimal("0.005"):
            out["flags"].append(f"{meat}: the spit weights add up to {parts} but the sheet's own total says {printed}")
        if not out["spits"][meat]:
            out["flags"].append(f"{meat}: no spit weight was found in the tab")
    for r in range(1, 20):
        if isinstance(ws.cell(r, 3).value, str) and "قبل الشك" in ws.cell(r, 3).value:
            out["opening"] = _weigh_row(ws, r)
            break
    lab = _find(ws, "end-of-shift closing summary")
    if lab:
        for r in range(lab[0].row, lab[0].row + 12):
            if any(isinstance(ws.cell(r, c).value, str) and "قبل الشك" in ws.cell(r, c).value for c in (3, 14)):
                out["closing"] = _weigh_row(ws, r)
                break
    if out["closing"] is None:
        out["flags"].append("no closing weigh-in (the end-of-shift summary) was found in the tab")
    return out


def _total(w):
    return sum((v or Decimal(0)) for v in w.values()) if w else None


def fridge_check(day_read, received, waste=None):
    """Closing weigh-in against opening + received - spits - waste (all raw kg; trim + main), per meat. Informational: measured on the 35 real days the gap runs
    from -4 to +4 kg either way (bins, trimmings, the baselines of Sep 1-2 and two odd days reach 7-11 kg), so it is a note for the manager's eye, never a gate.
    Returns {meat: {"expected_kg", "closing_kg", "gap_kg" (closing - expected), "large" (more than 5 kg)}} for the meats it can compute."""
    out = {}
    op, cl = day_read.get("opening"), day_read.get("closing")
    if not op or not cl:
        return out
    for meat in ("chicken", "beef"):
        rec = received.get(meat)
        if rec is None or not day_read["spits"][meat]:
            continue
        spits = sum(day_read["spits"][meat], Decimal(0))
        o = (op[meat].get("trim_kg") or 0) + (op[meat].get("main_kg") or 0)
        c = (cl[meat].get("trim_kg") or 0) + (cl[meat].get("main_kg") or 0)
        wst = Decimal(str((waste or {}).get(meat) or 0))
        expected = Decimal(str(o)) + Decimal(str(rec)) - spits - wst
        gap = Decimal(str(c)) - expected
        out[meat] = {"expected_kg": expected, "closing_kg": Decimal(str(c)), "gap_kg": gap, "large": abs(gap) > FRIDGE_LARGE_KG}
    return out


def meat_weights_for_approval(day, received_from_ledger=None, waste_from_ledger=None, restaurant_root=None, extra_dirs=None):
    """The `meat_weights` block of an SQL approval, from the workbook: the spits (their sum) and the closing weigh-in per meat. Received is NOT in the workbook:
    it is passed through from the company ledger (a delivery is a number only the manager knows; the ledger was wrong once, Sep 23) and the notes say so.
    Returns (meat_weights, notes, flags) -- or (None, [], [reason]) when there is no workbook."""
    path, why = find_workbook(day, restaurant_root, extra_dirs)
    if path is None:
        return None, [], [why]
    r = read_day(path, day)
    notes, flags = [f"weights from {why}, tab {r['sheet']}"], list(r["flags"])
    out = {}
    for meat in ("chicken", "beef"):
        w = {}
        if r["spits"][meat]:
            w["spit_kg"] = float(sum(r["spits"][meat], Decimal(0)))
        if r["closing"]:
            c = r["closing"][meat]
            if c.get("trim_kg") is not None and c.get("main_kg") is not None:
                w["closing"] = {"trim_kg": float(c["trim_kg"]), "main_kg": float(c["main_kg"]), "tallow_kg": float(c.get("tallow_kg") or 0)}
        if received_from_ledger and received_from_ledger.get(meat) is not None:
            w["received_kg"] = float(received_from_ledger[meat])
            notes.append(f"{meat} RECEIVED {received_from_ledger[meat]} kg comes from the company ledger, not from the workbook: confirm the delivery")
        out[meat] = w
    for meat, f in fridge_check(r, received_from_ledger or {}, waste_from_ledger).items():
        line = f"{meat} fridge: opening + received - spits - waste = {f['expected_kg']} kg against the closing weigh-in {f['closing_kg']} kg (gap {f['gap_kg']:+} kg)"
        (flags if f["large"] else notes).append(line + (": unusually large, check the delivery or the weigh-in" if f["large"] else ""))
    return out, notes, flags


def main(argv=None):
    if hasattr(sys.stdout, "reconfigure"):
        try:
            sys.stdout.reconfigure(encoding="utf-8")
        except Exception:
            pass
    argv = list(sys.argv[1:] if argv is None else argv)
    if not argv or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", argv[0]):
        print("usage: py restaurant_weights.py YYYY-MM-DD")
        return 2
    path, why = find_workbook(argv[0])
    print(why)
    if path is None:
        return 1
    r = read_day(path, argv[0])
    print("tab:", r["sheet"])
    for m in ("chicken", "beef"):
        print(f"{m}: spits {[str(x) for x in r['spits'][m]]} = {sum(r['spits'][m], Decimal(0))}; sheet total {r['spit_sum_cell'].get(m)}")
    print("opening:", r["opening"])
    print("closing:", r["closing"])
    for f in r["flags"]:
        print("FLAG:", f)
    return 0


if __name__ == "__main__":
    sys.exit(main())
