"""
test_restaurant_api.py -- offline tests for the HUD's Restaurant page backend (ui/restaurant_api.py), through FastAPI's TestClient.

The Demo Bistro tree, the app backup, the inbox and the output folders are temp folders; the photo reading is replaced by the saved demo draft (examples/demo_bistro).
Nothing touches the real folders, Supabase or Ollama.   Run: py -3.11 test_restaurant_api.py
"""
from david_config import DAVID_DATA, DAVID_HOME
import sys, os, json, glob, re, copy, shutil, tempfile, inspect, time, threading
sys.stdout.reconfigure(encoding="utf-8")
sys.path.insert(0, str(DAVID_HOME))
sys.path.insert(0, str(DAVID_HOME / "ui"))

PASS = "\033[92m  PASS\033[0m"
FAIL = "\033[91m  FAIL\033[0m"
results = []


def check(name, ok, detail=""):
    results.append(bool(ok))
    print(f"{PASS if ok else FAIL}  {name}" + (f" -- {detail}" if detail else ""))


import warnings
warnings.filterwarnings("ignore")
import openpyxl
from fastapi import FastAPI
from starlette.testclient import TestClient
import restaurant_pipeline as P
import restaurant_sql as S
import restaurant_api as A
import restaurant_corrections as RC

TMP = tempfile.mkdtemp(prefix="api_test_")
DAY = "2026-09-22"
ROOT = os.path.join(TMP, "abo")
os.makedirs(os.path.join(ROOT, "app"))
open(os.path.join(ROOT, "app", ".env.production"), "w").write("VITE_SUPABASE_URL=https://abcdefghijklmnop.supabase.co\n")
INBOX, DAYS, REV = os.path.join(TMP, "inbox"), os.path.join(TMP, "days"), os.path.join(TMP, "review")
BK = os.path.join(ROOT, "app", "backup", "2026-09-22")
os.makedirs(BK)
for t in S.TABLES:
    json.dump([], open(os.path.join(BK, t + ".json"), "w", encoding="utf-8"))

gate = threading.Event()
gate.set()
REAL = json.load(open(str(DAVID_HOME / "examples" / "demo_bistro" / "draft_package.demo.json"), encoding="utf-8"))   # the demo draft package (synthetic)
runs = []


def fake_run_day(folder, *, business_day=None, drafts_dir=None, log=print, **kw):
    runs.append(folder)
    log("[1/5] reading the photos")
    gate.wait(10)
    pkg = copy.deepcopy(REAL)
    nn = len(glob.glob(os.path.join(drafts_dir, f"DRAFT-{business_day}-*.json"))) + 1
    pkg["draft_id"] = f"DRAFT-{business_day}-{nn:03d}"
    path = os.path.join(drafts_dir, pkg["draft_id"] + ".json")
    json.dump(pkg, open(path, "w", encoding="utf-8"), ensure_ascii=False)
    return {"ok": True, "package": pkg, "package_path": path, "validation": {"valid": True, "errors": []}, "review_path": None}


P.run_day = fake_run_day


def book_bytes(path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("22-09-2026")
    ws.cell(3, 3).value = "قبل الشك"
    for c, v in ((4, 4.0), (5, 150.0), (15, 1.0), (16, 6.0), (17, 28.0)):
        ws.cell(3, c).value = v
    ws.cell(6, 2).value = "الوزن السيخ"; ws.cell(6, 4).value = 35.68
    ws.cell(6, 14).value = "الوزن السيخ"; ws.cell(6, 17).value = 7.7
    ws.cell(20, 2).value = "مجموع السياخ دجاج"; ws.cell(22, 2).value = 35.68
    ws.cell(20, 14).value = "مجموع السياخ لحم"; ws.cell(22, 14).value = 7.7
    ws.cell(30, 1).value = "end-of-shift closing summary"
    ws.cell(31, 3).value = "قبل الشك"
    for c, v in ((4, 5.0), (5, 182.0), (15, 1.5), (16, 6.0), (17, 28.5)):
        ws.cell(31, c).value = v
    wb.save(path)
    return open(path, "rb").read()


XLSX = book_bytes(os.path.join(TMP, "w.xlsx"))
app = FastAPI()
WATCHD = os.path.join(TMP, "watch")
app.include_router(A.make_router({"days_dir": DAYS, "review_dir": REV, "inbox_dir": INBOX, "restaurant_root": ROOT, "backup": BK, "watch_dir": WATCHD}))
c = TestClient(app, base_url="http://127.0.0.1:8000")


def up(day, name, data=b"x" * 100):
    return c.put("/api/restaurant/upload", params={"day": day, "name": name}, content=data)


def wait_job(timeout=20):
    t0 = time.time()
    while time.time() - t0 < timeout:
        j = c.get("/api/restaurant/job").json()
        if j.get("state") in ("done", "error"):
            return j
        time.sleep(0.1)
    return c.get("/api/restaurant/job").json()


# ── overview ─────────────────────────────────────────────────────────────────────────────────────────────────────────────────────
o = c.get("/api/restaurant/overview")
check("the overview works on empty folders: no days, a default day, the newest backup named", o.status_code == 200 and o.json()["days"] == [] and o.json()["backup"] == "2026-09-22" and o.json()["job"] == {"state": "idle"})
check("with an empty app the default day is today (nothing to continue from)", re.fullmatch(r"\d{4}-\d{2}-\d{2}", o.json()["default_day"]) is not None)

# ── upload: a folder's files go into David's own inbox, never anywhere else ──────────────────────────────────────────────────
r = up(DAY, "WhatsApp Image 2026-09-23 at 2.15.07 AM.jpeg")
check("a photo is stored in the inbox of that day, under its own name", r.status_code == 200 and r.json()["name"] == "WhatsApp Image 2026-09-23 at 2.15.07 AM.jpeg" and os.path.isfile(os.path.join(INBOX, DAY, "WhatsApp Image 2026-09-23 at 2.15.07 AM.jpeg")))
r = up(DAY, "WhatsApp Image 2026-09-23 at 2.15.07 AM.jpeg")
check("the same name twice is kept twice (the second gets (2)), nothing is overwritten", r.json()["name"] == "WhatsApp Image 2026-09-23 at 2.15.07 AM (2).jpeg" and len(os.listdir(os.path.join(INBOX, DAY))) == 2)
r = up(DAY, "..\\..\\..\\Windows\\evil.jpg")
check("a path in the file name is cut down to the file name, so nothing can leave the inbox", r.status_code == 200 and r.json()["name"] == "evil.jpg" and not os.path.exists(os.path.join(TMP, "Windows")) and os.path.isfile(os.path.join(INBOX, DAY, "evil.jpg")))
r = up(DAY, "../../x.jpg")
check("a forward-slash path is cut down the same way", r.json()["name"] == "x.jpg")
check("a program or script is refused (415), and nothing is written for it", up(DAY, "run.exe").status_code == 415 and up(DAY, "x.bat").status_code == 415 and up(DAY, "noext").status_code == 415 and not os.path.exists(os.path.join(INBOX, DAY, "run.exe")))
check("an empty file is refused and leaves no part file behind", up(DAY, "empty.jpg", b"").status_code == 400 and not any(f.endswith(".part") for f in os.listdir(os.path.join(INBOX, DAY))) and not os.path.exists(os.path.join(INBOX, DAY, "empty.jpg")))
check("a day that is not a date is refused (400) before anything is written", up("../etc", "a.jpg").status_code == 400 and up("2026-13-01", "a.jpg").status_code == 400 and sorted(os.listdir(INBOX)) == [DAY])
orig_max = A.MAX_FILE_BYTES
A.MAX_FILE_BYTES = 50
big = up(DAY, "big.jpg", b"y" * 200)
A.MAX_FILE_BYTES = orig_max
check("a file over the size limit is refused (413) and its part file is removed", big.status_code == 413 and not os.path.exists(os.path.join(INBOX, DAY, "big.jpg")) and not any(f.endswith(".part") for f in os.listdir(os.path.join(INBOX, DAY))))
orig_day = A.MAX_DAY_BYTES
A.MAX_DAY_BYTES = 300
over = up(DAY, "more.jpg", b"z" * 400)
A.MAX_DAY_BYTES = orig_day
check("a day's total over the limit is refused (413)", over.status_code == 413)
for f in os.listdir(os.path.join(INBOX, DAY)):
    os.remove(os.path.join(INBOX, DAY, f))
check("the Demo Bistro tree is never written: no file in it after all of that", not any(os.path.isfile(p) for p in glob.glob(os.path.join(ROOT, "**", "*.jpg"), recursive=True)))

# ── the recent day folders (read-only status for the page's card) ────────────────────────────────────────────────────────────────
import datetime as _dt, time as _time
_today = _dt.date.today()


def _mkday(d, photos, sheets, age=1000):
    f = os.path.join(ROOT, f"{d.month:02d}-{d.year}", f"{d.day:02d}-{d.month:02d}-{d.year}")
    os.makedirs(f, exist_ok=True)
    for i in range(photos):
        open(os.path.join(f, f"WhatsApp Image {i}.jpeg"), "wb").write(b"x" * 100)
    for i in range(sheets):
        open(os.path.join(f, f"sheet {i} {d.day:02d}-{d.month:02d}-{d.year}.xlsx"), "wb").write(b"y" * 100)
    for n in os.listdir(f):
        _t = _time.time() - age
        os.utime(os.path.join(f, n), (_t, _t))
    return f


_ready_dir = _mkday(_today, 5, 3)
_wait_dir = _mkday(_today - _dt.timedelta(days=1), 2, 1)
_arr_dir = _mkday(_today - _dt.timedelta(days=2), 5, 3, age=5)
_before = sorted(os.listdir(_ready_dir))
fo = c.get("/api/restaurant/folders")
rows = {r["day"]: r for r in fo.json()["folders"]} if fo.status_code == 200 else {}
check("the folders card: a complete day (5 photos, 3 sheets) is 'ready', a day with 2 photos and 1 sheet is 'waiting' and says what is missing, a day whose files just changed is 'arriving'",
      fo.status_code == 200 and rows[_today.isoformat()]["state"] == "ready" and rows[(_today - _dt.timedelta(days=1)).isoformat()]["state"] == "waiting"
      and "2 more Excel sheet" in rows[(_today - _dt.timedelta(days=1)).isoformat()]["message"] and rows[(_today - _dt.timedelta(days=2)).isoformat()]["state"] == "arriving", str(rows)[:400])
check("it reports the counts and the need (4 photos, 3 sheets), reads only (the day folder is unchanged), and cannot be written to",
      rows[_today.isoformat()]["photos"] == 5 and rows[_today.isoformat()]["sheets"] == 3 and fo.json()["need"] == {"photos": 4, "sheets": 3} and sorted(os.listdir(_ready_dir)) == _before and c.post("/api/restaurant/folders", json={}).status_code == 405)
_dd = os.path.join(DAYS, _today.isoformat()); os.makedirs(_dd, exist_ok=True)
open(os.path.join(_dd, "DRAFT-%s-001.json" % _today.isoformat()), "w").write("{}")
check("a day that has a draft shows as 'read'; with an approval 'approved'; with an upload 'uploaded' (and days older than the newest uploaded one disappear: they are history)",
      {r["day"]: r["state"] for r in c.get("/api/restaurant/folders").json()["folders"]}[_today.isoformat()] == "read"
      and (open(os.path.join(_dd, "approval.json"), "w").write("{}") or True) and {r["day"]: r["state"] for r in c.get("/api/restaurant/folders").json()["folders"]}[_today.isoformat()] == "approved"
      and (open(os.path.join(_dd, "upload.json"), "w").write("{}") or True)
      and {r["day"]: r["state"] for r in c.get("/api/restaurant/folders").json()["folders"]} == {_today.isoformat(): "uploaded"})
shutil.rmtree(_dd, ignore_errors=True)
for _f in (_ready_dir, _wait_dir, _arr_dir):
    shutil.rmtree(_f, ignore_errors=True)
    try:
        os.rmdir(os.path.dirname(_f))
    except OSError:
        pass

# ── the day-folder watcher's status (read-only) ─────────────────────────────────────────────────────────────────────────────────────
wz = c.get("/api/restaurant/watch")
check("the watcher status works before it ever ran", wz.status_code == 200 and wz.json()["days"] == {})
os.makedirs(WATCHD, exist_ok=True)
open(os.path.join(WATCHD, "status.json"), "w", encoding="utf-8").write(json.dumps({"updated": "x", "days": {"2026-10-12": {"state": "ready", "message": "m"}}}))
check("it shows what the watcher saved", c.get("/api/restaurant/watch").json()["days"]["2026-10-12"]["state"] == "ready")
check("it cannot be written to (GET only)", c.post("/api/restaurant/watch", json={}).status_code == 405)

# ── prepare as a background job ──────────────────────────────────────────────────────────────────────────────────────────────────
r = c.post("/api/restaurant/prepare", json={"day": DAY})
check("preparing a day with no files says so (422), not a crash", r.status_code == 422 and "no folder" in r.json()["detail"])
up(DAY, "WhatsApp Image 2026-09-23 at 2.15.07 AM.jpeg")
up(DAY, "فوارق الشك مع الاوزان 22-09-2026.xlsx", XLSX)
check("a day's files are listed with their sizes", [f["name"] for f in c.get(f"/api/restaurant/day/{DAY}").json()["inbox"]] == ["WhatsApp Image 2026-09-23 at 2.15.07 AM.jpeg", "فوارق الشك مع الاوزان 22-09-2026.xlsx"])
gate.clear()
r1 = c.post("/api/restaurant/prepare", json={"day": DAY})
time.sleep(0.4)
r2 = c.post("/api/restaurant/prepare", json={"day": DAY})
check("preparing starts a background job, and a second request while it runs is refused (409): one reading at a time, the graphics card is shared", r1.status_code == 200 and r1.json()["state"] == "running" and r2.status_code == 409)
check("the job shows what it is doing", "reading the photos" in c.get("/api/restaurant/job").json()["message"])
gate.set()
j = wait_job()
check("the job finishes as done, with the day and a message", j["state"] == "done" and j["day"] == DAY and j["message"] == "the draft is ready", str(j))
check("the draft was read from the INBOX folder, not from the Demo Bistro tree", runs and runs[-1] == os.path.join(INBOX, DAY))
d = c.get(f"/api/restaurant/day/{DAY}").json()
check("the day now has its questions (figures, deliveries, expenses, warnings) and is at step 2", d["step"] == 2 and d["source"] == "inbox" and d["questions"]["figures"]["sales"] == REAL["cash"]["pos_sales"] and set(d["questions"]["deliveries"]) == {"chicken", "beef"})
check("the weights workbook uploaded with the photos was found (the questions carry its notes)", any("RECEIVED" in n for n in d["questions"]["notes"] + d["questions"]["warnings"]))
ov = c.get("/api/restaurant/overview").json()
check("the overview lists the day with its step and its upload count", ov["days"][0]["day"] == DAY and ov["days"][0]["step"] == 2 and ov["days"][0]["inbox_files"] == 2 and ov["days"][0]["in_app"] is False)

# a failing job is reported, not hidden
orig_run = P.run_day
P.run_day = lambda *a, **k: (_ for _ in ()).throw(RuntimeError("model is not reachable"))
c.post("/api/restaurant/prepare", json={"day": DAY})
jf = wait_job()
P.run_day = orig_run
check("a job that fails says why, and the next job can start", jf["state"] == "error" and "model is not reachable" in jf["message"] and c.post("/api/restaurant/prepare", json={"day": DAY}).status_code == 200)
wait_job()

# ── answers ──────────────────────────────────────────────────────────────────────────────────────────────────────────────────────
check("answers with a wrong kind of value are refused (400)", c.put(f"/api/restaurant/day/{DAY}/answers", json={"deliveries_confirmed": "yes"}).status_code == 400
      and c.put(f"/api/restaurant/day/{DAY}/answers", json={"allow_updates": ["users"]}).status_code == 400
      and c.put(f"/api/restaurant/day/{DAY}/answers", json={"paid_from_safe": "gaz"}).status_code == 400
      and c.put(f"/api/restaurant/day/{DAY}/answers", json={"approved_by": "x" * 200}).status_code == 400
      and c.put(f"/api/restaurant/day/{DAY}/answers", json=[1]).status_code == 400)
r = c.put(f"/api/restaurant/day/{DAY}/answers", json={"approved_by": "Demo User", "evil_key": "x", "paid_from_safe": []})
saved = json.load(open(os.path.join(REV, f"day_{DAY}", "answers.json"), encoding="utf-8"))
check("only the known answer fields are kept; the file lands next to the day's other files", r.status_code == 200 and "evil_key" not in saved and saved["approved_by"] == "Demo User")
f0 = c.post(f"/api/restaurant/day/{DAY}/finish")
check("finish with answers missing returns 422 with the open items listed, and makes no files", f0.status_code == 422 and f0.json()["status"] == "answers_missing" and any("still to confirm or correct" in e for e in f0.json()["errors"]) and not os.path.exists(os.path.join(REV, f"day_{DAY}", "apply.sql")))
q = d["questions"]
ALLC = {k: True for k, _ in RC.required_items(q)}
check("confirming items is saved as the page does it (item by item), and a confirmation must be true", c.put(f"/api/restaurant/day/{DAY}/answers", json={"confirmed": {"cash.pos_sales": False}}).status_code == 400
      and c.put(f"/api/restaurant/day/{DAY}/answers", json={"confirmed": ALLC, "cleared_blockers": {str(i): "looked at it" for i in range(1, len(q["blockers"]) + 1)}}).status_code == 200)
fx = c.put(f"/api/restaurant/day/{DAY}/answers", json={"overrides": {"cash": {"pos_sales": -1}}})
check("a correction that is not valid is refused (400) with the reason, a good one is saved", fx.status_code == 400 and "not valid" in fx.json()["detail"]
      and c.put(f"/api/restaurant/day/{DAY}/answers", json={"overrides": {"cash": {"pos_sales": REAL["cash"]["pos_sales"] + 3}}}).status_code == 200
      and c.get(f"/api/restaurant/day/{DAY}").json()["answers"]["overrides"] == {"cash": {"pos_sales": REAL["cash"]["pos_sales"] + 3.0}})
c.put(f"/api/restaurant/day/{DAY}/answers", json={"overrides": {}})
check("notes are saved per item, and a note that is too long, not text, or a list is refused (400)", c.put(f"/api/restaurant/day/{DAY}/answers", json={"notes": {"cash.pos_sales": "looked at the receipt"}}).status_code == 200
      and c.get(f"/api/restaurant/day/{DAY}").json()["answers"]["notes"] == {"cash.pos_sales": "looked at the receipt"}
      and c.put(f"/api/restaurant/day/{DAY}/answers", json={"notes": {"a": "x" * 301}}).status_code == 400
      and c.put(f"/api/restaurant/day/{DAY}/answers", json={"notes": {"a": 5}}).status_code == 400
      and c.put(f"/api/restaurant/day/{DAY}/answers", json={"notes": ["x"]}).status_code == 400)
c.put(f"/api/restaurant/day/{DAY}/answers", json={"notes": {}})
menu = c.get("/api/restaurant/menu").json()["items"]
check("the menu list for the 'which item is it' choice comes from the sheet: ids and Arabic names, the chicken combo among them", any(i["id"] == "C20" and i["name_ar"] for i in menu) and len(menu) > 70)
f1 = c.post(f"/api/restaurant/day/{DAY}/finish")
check("with the answers in, finish makes the SQL, the rehearsal passed, and the files are listed", f1.status_code == 200 and f1.json()["ok"] and f1.json()["status"] == "sql_ready" and f1.json()["rehearsal"]["ok"] and f1.json()["files"]["apply"] and f1.json()["files"]["undo"] and f1.json()["files"]["steps"], str(f1.json())[:300])
ap = c.get(f"/api/restaurant/day/{DAY}/file/apply").json()
check("apply.sql can be read through the page (and copied): begin, inserts, commit, the check table", ap["text"].lstrip().startswith("-- apply.sql for") and "\nbegin;\n" in ap["text"] and "\ncommit;\n" in ap["text"] and "rows_found" in ap["text"])
st = c.get(f"/api/restaurant/day/{DAY}/file/steps").json()["text"]
check("the step-by-step file has the SQL editor address and the export command with a placeholder, never a key", "supabase.com/dashboard/project/abcdefghijklmnop/sql/new" in st and "PASTE_KEY_HERE" in st)
check("an unknown file kind is 404, a file not made yet is 404, a path in the kind cannot reach another file", c.get(f"/api/restaurant/day/{DAY}/file/passwd").status_code == 404 and c.get(f"/api/restaurant/day/{DAY}/file/check").status_code == 404 and c.get(f"/api/restaurant/day/{DAY}/file/..%2F..%2Fanswers").status_code == 404)
check("the day is now at step 3 (run the SQL)", c.get(f"/api/restaurant/day/{DAY}").json()["step"] == 3)
ck = c.post(f"/api/restaurant/day/{DAY}/check").json()
check("check before the SQL was run says the day is not in the backup yet", ck["status"] == "not_in_backup")

# ── clearing an inbox ────────────────────────────────────────────────────────────────────────────────────────────────────────────
rm = c.delete(f"/api/restaurant/inbox/{DAY}")
check("clearing a day's inbox removes only the uploaded copies (the drafts and the SQL files stay)", rm.json()["removed"] == 2 and not os.path.exists(os.path.join(INBOX, DAY)) and os.path.isfile(os.path.join(REV, f"day_{DAY}", "apply.sql")))
check("clearing an inbox that is not there is not an error", c.delete("/api/restaurant/inbox/2026-01-01").json()["removed"] == 0)

# ── the real HUD app: the router is mounted, host and origin guards cover it ──────────────────────────────────────────────────
import server as srv
hud = TestClient(srv.app, base_url="http://127.0.0.1:8000")
check("the HUD serves the Restaurant API (the router is mounted before the page fallback)", hud.get("/api/restaurant/overview").status_code == 200 and "days" in hud.get("/api/restaurant/overview").json())
check("a page on another site cannot POST to it (the HUD's origin guard answers 403)", hud.post("/api/restaurant/prepare", json={"day": "2026-10-08"}, headers={"Origin": "http://evil.example"}).status_code == 403
      and hud.put("/api/restaurant/upload", params={"day": "2026-10-08", "name": "a.jpg"}, content=b"x", headers={"Origin": "http://evil.example"}).status_code == 403)
check("a foreign Host header is refused (DNS-rebinding guard)", TestClient(srv.app, base_url="http://evil.example").get("/api/restaurant/overview").status_code == 400)
check("an unknown path under /api/restaurant is a 404 JSON, not the web page", hud.get("/api/restaurant/nope").status_code in (404, 405))

code = "\n".join(l for l in inspect.getsource(A).splitlines() if not l.lstrip().startswith("#"))
check("the API never connects to anything: no http client, socket, supabase or database import", not re.search(r"(?m)^\s*(import|from)\s+(requests|httpx|socket|urllib|sqlite3|psycopg|supabase)", code))
check("it writes only two kinds of file (an upload into the inbox, and the answers file) and removes only the day's own inbox folder", len(re.findall(r'open\([^)]*"w"[,)]', code)) == 1 and code.count('open(tmp, "wb")') == 1 and code.count("shutil.rmtree(") == 1 and "realpath(box).startswith(os.path.realpath(inbox_dir))" in code)

shutil.rmtree(TMP, ignore_errors=True)
print(f"\n  {sum(results)}/{len(results)} passed")
sys.exit(0 if all(results) else 1)
