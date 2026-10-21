"""
test_restaurant_day.py -- offline tests for restaurant_day.py (the daily runner) and restaurant_rehearsal.py (the throw-away rehearsal of a generated SQL file).

No model, no network, no Supabase, no real folder is written: the Demo Bistro tree, the backup and the output folders are built in a temp directory, the photo reading is
replaced by the saved demo draft (examples/demo_bistro).   Run: py -3.11 test_restaurant_day.py
"""
from david_config import DAVID_DATA, DAVID_HOME
import sys, os, json, glob, re, copy, shutil, tempfile, inspect, datetime
sys.stdout.reconfigure(encoding="utf-8")
sys.path.insert(0, str(DAVID_HOME))

PASS = "\033[92m  PASS\033[0m"
FAIL = "\033[91m  FAIL\033[0m"
results = []


def check(name, ok, detail=""):
    results.append(bool(ok))
    print(f"{PASS if ok else FAIL}  {name}" + (f" -- {detail}" if detail else ""))


import openpyxl
import restaurant_day as D
import restaurant_rehearsal as RH
import restaurant_sql as S
import restaurant_pipeline as P

TMP = tempfile.mkdtemp(prefix="day_test_")
DAY = "2026-09-22"
BR = "00000000-0000-0000-0000-000000000001"

# ── the fake Demo Bistro tree: the day's folder, a weights workbook, the app's env file ────────────────────────────────────────────
ROOT = os.path.join(TMP, "abo")
FOLDER = os.path.join(ROOT, "09-2026", "22-09-2026")
os.makedirs(FOLDER)
os.makedirs(os.path.join(ROOT, "app"))
open(os.path.join(ROOT, "app", ".env.production"), "w").write("VITE_SUPABASE_URL=https://abcdefghijklmnop.supabase.co\nVITE_SUPABASE_ANON_KEY=SECRET-SHOULD-NEVER-BE-READ\n")


def make_book(path, tab, chicken=(35.68,), beef=(7.7,)):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet(tab)
    ws.cell(3, 3).value = "قبل الشك"
    for c, v in ((4, 4.0), (5, 150.0), (15, 1.0), (16, 6.0), (17, 28.0)):
        ws.cell(3, c).value = v
    for i, w in enumerate(chicken):
        ws.cell(6 + i, 2).value = "الوزن السيخ"
        ws.cell(6 + i, 4).value = w
    for i, w in enumerate(beef):
        ws.cell(6 + i, 14).value = "الوزن السيخ"
        ws.cell(6 + i, 17).value = w
    ws.cell(20, 2).value = "مجموع السياخ دجاج"
    ws.cell(22, 2).value = round(sum(chicken), 3)
    ws.cell(20, 14).value = "مجموع السياخ لحم"
    ws.cell(22, 14).value = round(sum(beef), 3)
    ws.cell(30, 1).value = "end-of-shift closing summary"
    ws.cell(31, 3).value = "قبل الشك"
    for c, v in ((4, 5.0), (5, 182.0), (15, 1.5), (16, 6.0), (17, 28.5)):
        ws.cell(31, c).value = v
    wb.save(path)


make_book(os.path.join(FOLDER, "فوارق الشك مع الاوزان 22-09-2026.xlsx"), "22-09-2026")

# ── the saved demo draft stands in for the photo reading ─────────────────────────────────────────────────────────────────────────
real_path = str(DAVID_HOME / "examples" / "demo_bistro" / "draft_package.demo.json")   # the demo draft package (synthetic)
REAL = json.load(open(real_path, encoding="utf-8"))
REAL_VALID = {"valid": True, "errors": []}
HELD = REAL["extraction"]["summary_by_items"]["lines"][0]
HELD.update(readings_agree=False, name_ar_alt="سعتنتز", qty_alt=None)
calls = []


def fake_run_day(folder, *, business_day=None, drafts_dir=None, **kw):
    calls.append((folder, business_day, kw.get("day_status"), kw.get("correction_reason")))
    pkg = copy.deepcopy(REAL)
    nn = len(glob.glob(os.path.join(drafts_dir, f"DRAFT-{business_day}-*.json"))) + 1
    pkg["draft_id"] = f"DRAFT-{business_day}-{nn:03d}"
    path = os.path.join(drafts_dir, pkg["draft_id"] + ".json")
    json.dump(pkg, open(path, "w", encoding="utf-8"), ensure_ascii=False)
    return {"ok": True, "package": pkg, "package_path": path, "validation": REAL_VALID, "review_path": None}


P.run_day = fake_run_day


def empty_backup_dir(path, days=None, extra=None):
    os.makedirs(path, exist_ok=True)
    for t in S.TABLES:
        json.dump((extra or {}).get(t, []), open(os.path.join(path, t + ".json"), "w", encoding="utf-8"), ensure_ascii=False)
    if days:
        json.dump(days, open(os.path.join(path, "days.json"), "w", encoding="utf-8"))
    return path


BK0 = empty_backup_dir(os.path.join(ROOT, "app", "backup", "2026-09-22"))
DAYS, REV = os.path.join(TMP, "days"), os.path.join(TMP, "review")
kw = dict(restaurant_root=ROOT, backup=BK0, days_dir=DAYS, review_dir=REV)

# ── prepare ────────────────────────────────────────────────────────────────────────────────────────────────────────────────────────
r = D.prepare(DAY, log=lambda m: None, **kw)
check("prepare reads the day's folder as an ENDED day and saves the draft under the day's own folder", r["ok"] and calls[-1][0].endswith("22-09-2026") and calls[-1][2] == "ended" and r["draft_path"].startswith(os.path.join(DAYS, DAY)) and os.path.isfile(r["draft_path"]))
q = r["questions"]
check("the questions carry the figures, the deliveries, the expenses and the draft's own flags", q["figures"]["sales"] == REAL["cash"]["pos_sales"] and set(q["deliveries"]) == {"chicken", "beef"}
      and [e["label"] for e in q["expenses"]] == [e["label"] for e in REAL["cash"]["expenses"]] and q["checks"] == [f for f in REAL["flags"] if f.startswith("CHECK:")] and q["blockers"] == [f for f in REAL["flags"] if f.startswith("BLOCKER:")])
md = open(os.path.join(r["pack_dir"], "questions.md"), encoding="utf-8").read()
check("questions.md names the deliveries, the expenses to place in the safe, and every warning, and shows item names, not codes", "Deliveries" in md and "paid from the SAFE" in md and all(c[:60] in md for c in q["checks"]) and not re.search(r"\b[CKXU]\d{2}\b", md.split("## Questions")[1]))
check("the weights workbook's notes are in the questions (received comes from the ledger; the fridge line)", any("RECEIVED" in n for n in q["notes"] + q["warnings"]) and any("fridge" in n for n in q["notes"] + q["warnings"]))
check("the chain is reported: with no day before in the backup it says the opening could not be checked", any("not in the newest backup" in n for n in q["notes"]))
tpl = json.load(open(r["answers_path"], encoding="utf-8"))
check("an answers.json template is written with nothing approved and nothing confirmed", tpl["approved_by"] == "" and tpl["deliveries_confirmed"] is False and tpl["acknowledge_checks"] is False and set(tpl["cleared_blockers"]) == {str(i) for i in range(1, len(q["blockers"]) + 1)})
tpl["approved_by"] = "kept"
json.dump(tpl, open(r["answers_path"], "w", encoding="utf-8"))
D.prepare(DAY, log=lambda m: None, **kw)
check("running prepare again makes a NEW draft number and never overwrites the manager's answers file", len(glob.glob(os.path.join(DAYS, DAY, f"DRAFT-{DAY}-*.json"))) == 2 and json.load(open(r["answers_path"], encoding="utf-8"))["approved_by"] == "kept")
tpl["approved_by"] = ""
json.dump(tpl, open(r["answers_path"], "w", encoding="utf-8"))
try:
    D.prepare("2026-02-31", log=lambda m: None, **kw)
    ok = False
except D.DayError:
    ok = True
check("a date that is not a date is refused", ok)
try:
    D.prepare("2026-01-05", log=lambda m: None, **kw)
    ok = False
except D.DayError as e:
    ok = "no folder" in str(e)
check("a day with no folder is refused, naming the problem", ok)

# ── finish: the answers are the gate ───────────────────────────────────────────────────────────────────────────────────────────────
try:
    D.finish(DAY, **kw)
    ok = False
except D.DayError as e:
    msg = str(e)
    ok = "approved_by is empty" in msg and "still to confirm or correct" in msg and ("Warning:" in msg or not q["checks"]) and (("BLOCKER 1 has no reason" in msg) or not q["blockers"])
check("with the template as it is (nothing approved) finish stops and lists every missing answer at once", ok, msg if not ok else "")
good = {"approved_by": "Demo User (answered in chat)", "deliveries_confirmed": True, "acknowledge_checks": True, "paid_from_safe": [], "cleared_blockers": {str(i): "checked against the Helper" for i in range(1, len(q["blockers"]) + 1)}, "allow_updates": [], "held_lines_all_correct": True}
f = D.finish(DAY, answers=good, **kw)
check("with real answers finish produces SQL and the rehearsal passes", f["ok"] and f["status"] == "sql_ready" and f["rehearsal"]["ok"] and f["rehearsal"]["checked"] >= 12, str(f.get("reasons")) + str(f.get("rehearsal")))
names = sorted(os.path.basename(x) for x in f["files"])
check("the manager's folder gets apply.sql, undo.sql, plan.md and RUN_STEPS.md", names == ["RUN_STEPS.md", "apply.sql", "plan.md", "undo.sql"] and all(os.path.isfile(x) for x in f["files"]))
steps = open(os.path.join(f["pack_dir"], "RUN_STEPS.md"), encoding="utf-8").read()
check("RUN_STEPS has the SQL editor address of the project (read from the app's env, address only), the export command with a placeholder and the undo instruction",
      "https://supabase.com/dashboard/project/abcdefghijklmnop/sql/new" in steps and "PASTE_KEY_HERE" in steps and "export-backup.mjs" in steps and "undo.sql" in steps and f"py restaurant_day.py check {DAY}" in steps)
check("no key or secret value from the env file ever reaches an output", "SECRET-SHOULD-NEVER-BE-READ" not in steps and not any("SECRET-SHOULD" in open(x, encoding="utf-8").read() for x in f["files"]))
ap_saved = json.load(open(os.path.join(DAYS, DAY, "approval.json"), encoding="utf-8"))
check("the approval is saved next to the draft: who, the draft's fingerprint, the workbook figures", ap_saved["approved_by"] == good["approved_by"] and ap_saved["draft_id"] == D.latest_draft(DAY, DAYS).split(os.sep)[-1][:-5] and ap_saved["meat_weights"]["chicken"]["spit_kg"] == 35.68 and ap_saved["meat_weights"]["beef"]["spit_kg"] == 7.7)
check("held_lines_all_correct writes the manager's answers file for the learning step and names the command (it does NOT apply it)", "learn_command" in f and f["learn_command"].startswith("py restaurant_learn.py --answers") and os.path.isfile(os.path.join(DAYS, DAY, "manager_line_answers.json")))

# the rehearsal on a real generated file, standing alone
sql = open(os.path.join(f["pack_dir"], "apply.sql"), encoding="utf-8").read()
und = open(os.path.join(f["pack_dir"], "undo.sql"), encoding="utf-8").read()
reh = RH.rehearse(sql, und, BK0)
check("the rehearsal adds the day, the cash sheet, the expenses, the meat rows, and nothing else", reh["ok"] and reh["rows_added"]["days"] == 1 and reh["rows_added"]["cash_sheets"] == 1 and reh["rows_added"]["meat_closing_weights"] == 2 and set(reh["rows_added"]) <= set(RH.COLS), str(reh))
bad_sql = sql.replace("insert into cash_sheets", "insert into cash_sheetz", 1)
r2 = RH.rehearse(bad_sql, und, BK0)
check("a file that would fail is caught: a misspelt table gives a problem naming the error, not a pass", not r2["ok"] and any("failed on the rehearsal copy" in p for p in r2["problems"]), str(r2["problems"]))
r3 = RH.rehearse(sql.replace("1 as rows_expected", "2 as rows_expected", 1), und, BK0)
check("a check table whose line is not found is reported with its name", not r3["ok"] and any("not found after apply" in p for p in r3["problems"]))
r4 = RH.rehearse(sql, "begin;\ncommit;\n", BK0)
check("an undo that does not put the rows back is reported", not r4["ok"] and any("undo.sql did not put the rows back" in p for p in r4["problems"]))
r5 = RH.rehearse(sql, und, BK0)
check("the rehearsal never touches the backup it was given", json.load(open(os.path.join(BK0, "cash_sheets.json"), encoding="utf-8")) == [] and r5["ok"])

# a rehearsal failure stops the delivery
orig = RH.rehearse
RH.rehearse = lambda a, u, b: {"ok": False, "problems": ["forced for the test"], "checked": 0, "rows_added": {}}
fx = D.finish(DAY, answers=good, **kw)
RH.rehearse = orig
check("if the rehearsal fails no file is handed over and the reason is given", fx["ok"] is False and fx["status"] == "rehearsal_failed" and fx["reasons"] == ["forced for the test"])

# the other gates of the generator still apply
bad = dict(good, paid_from_safe=["not an expense"])
fb = D.finish(DAY, answers=bad, **kw)
check("an expense named as paid from the safe that is not in the draft stops it", fb["ok"] is False and fb["status"] == "refused" and any("not in the draft" in x for x in fb["reasons"]))
exp_label = REAL["cash"]["expenses"][0]["label"]
fs = D.finish(DAY, answers=dict(good, paid_from_safe=[exp_label]), **kw)
check("an expense the manager says came from the safe is written as the safe", fs["ok"] and f"'{exp_label}'" in open(os.path.join(fs["pack_dir"], "apply.sql"), encoding="utf-8").read() and "'safe'" in open(os.path.join(fs["pack_dir"], "apply.sql"), encoding="utf-8").read())
fs0 = D.finish(DAY, answers=good, **kw)

# ── check: after the manager ran it ────────────────────────────────────────────────────────────────────────────────────────────────
res = D.check(DAY, **kw)
check("check says the day is not in the backup yet when it is not (and says to export a fresh one)", res["status"] == "not_in_backup" and "export a fresh one" in res["message"])
con = RH.make_db(S.load_backup(BK0))
RH.run(con, open(os.path.join(f["pack_dir"], "apply.sql"), encoding="utf-8").read())
snap = RH.dump(con)
BK1 = os.path.join(ROOT, "app", "backup", "2026-09-23")
BOOL = ("closed", "float_open_ok", "float_close_ok")
empty_backup_dir(BK1, extra={t: [{k: (bool(v) if k in BOOL else v) for k, v in row.items() if v is not None} for row in snap[t]] for t in snap})
ck = D.check(DAY, **dict(kw, backup=BK1))
check("after the SQL ran, check says the app matches the draft (every row), and counts the fields it compared", ck["ok"] and ck["status"] == "matches" and ck["field_counts"].get("match", 0) >= 8 and ck["field_counts"].get("differs", 0) == 0, str(ck))
check("check writes check.md for the manager", "matches" in open(os.path.join(f["pack_dir"], "check.md"), encoding="utf-8").read())
cash = snap["cash_sheets"][0]
snap2 = copy.deepcopy(snap)
snap2["cash_sheets"][0]["sales_usd"] = snap2["cash_sheets"][0]["sales_usd"] + 0.5
BK2 = os.path.join(ROOT, "app", "backup", "2026-09-24")
empty_backup_dir(BK2, extra={t: [{k: (bool(v) if k in BOOL else v) for k, v in row.items() if v is not None} for row in snap2[t]] for t in snap2})
ck2 = D.check(DAY, **dict(kw, backup=BK2))
check("if the app differs from the draft (a sales figure typed 0.50 higher) check says so and shows the difference", ck2["ok"] is False and any("sales_usd" in d for d in ck2["differences"]) and ck2["field_counts"].get("differs", 0) >= 1, str(ck2))

# the day is already in the app: corrections need an explicit allow-update
fe = D.finish(DAY, answers=good, **dict(kw, backup=BK2))
check("when the day is already in the app and differs, finish stops (needs human) unless the answers allow updating that table", fe["ok"] is False and fe["status"] == "needs_human" and any("cash_sheets" in x for x in fe["needs_human"]))
fe2 = D.finish(DAY, answers=dict(good, allow_updates=["cash_sheets"]), **dict(kw, backup=BK2))
check("with allow_updates it makes a guarded correction and the rehearsal passes", fe2["ok"] and "update cash_sheets set sales_usd" in open(os.path.join(fe2["pack_dir"], "apply.sql"), encoding="utf-8").read())
fn = D.finish(DAY, answers=good, **dict(kw, backup=BK1))
check("when the app already holds exactly what the draft says there is nothing to do and no file is made", fn["ok"] and fn["status"] == "nothing_to_do")

# ── the manager's per-item answers: confirm or correct each item ───────────────────────────────────────────────────────────────────
import restaurant_corrections as RC
qq = D._questions(DAY, DAYS)
req = RC.required_items(qq)
allc = {k: True for k, _ in req}
base = {"approved_by": "Demo User", "confirmed": allc, "paid_from_safe": [], "cleared_blockers": {str(i): "looked" for i in range(1, len(qq["blockers"]) + 1)}, "allow_updates": [], "overrides": {}}
check("the questions list every figure as an item: 5 cash figures, each expense, 5 meat figures per meat, each warning, each held line",
      len(req) == 5 + len(qq["expenses"]) + 2 * (len(RC.MEAT_KEYS) + 1) + len(qq["checks"]) + len(qq["held_lines"]) and ("cash.pos_sales", "Sales") in req and ("meat.beef.closing", "Beef closing weights") in req and any(k.startswith("line.") for k, _ in req), str(len(req)))
try:
    D.finish(DAY, answers=dict(base, confirmed={}), **kw)
    ok = False
    m_ = ""
except D.DayError as e:
    m_ = str(e)
    ok = f"still to confirm or correct ({len(req)})" in m_ and "Sales" in m_ and "Chicken received" in m_
check("with nothing confirmed finish names every item still open, by its label", ok, m_ if not ok else "")
some = dict(allc)
some.pop("cash.pos_sales")
try:
    D.finish(DAY, answers=dict(base, confirmed=some), **kw)
    ok = False
except D.DayError as e:
    ok = "still to confirm or correct (1): Sales" in str(e)
check("one item left open stops it and says which", ok)
r_ok = D.finish(DAY, answers=base, **kw)
check("with every item confirmed finish makes the SQL (the per-item way replaces the single all-ok box)", r_ok["ok"] and r_ok["status"] == "sql_ready" and r_ok["rehearsal"]["ok"], str(r_ok.get("reasons")) + str(r_ok.get("status")))
r_fix = D.finish(DAY, answers=dict(base, confirmed=some, overrides={"cash": {"pos_sales": REAL["cash"]["pos_sales"] + 10}}), **kw)
apply_fix = open(os.path.join(r_fix["pack_dir"], "apply.sql"), encoding="utf-8").read()
exp_old = REAL["computed"]["expected"]
check("a corrected figure counts as settled, lands in the SQL, and Expected is recomputed from it (sales +10 -> Expected +10)",
      r_fix["ok"] and f"{REAL['cash']['pos_sales'] + 10:g}" in apply_fix and f"{round(exp_old + 10, 2):g}" in apply_fix, apply_fix[apply_fix.index("insert into cash_sheets"):][:420])
cd = json.load(open(os.path.join(DAYS, DAY, "corrected_draft.json"), encoding="utf-8"))
check("the corrected copy is saved (the saved draft is never edited), its cash and computed block agree, and the draft says it was corrected",
      cd["cash"]["pos_sales"] == REAL["cash"]["pos_sales"] + 10 and cd["computed"]["expected"] == round(exp_old + 10, 2) and any("manager corrected figures" in f for f in cd["flags"])
      and json.load(open(D.latest_draft(DAY, DAYS), encoding="utf-8"))["cash"]["pos_sales"] == REAL["cash"]["pos_sales"])
plan_fix = open(os.path.join(r_fix["pack_dir"], "plan.md"), encoding="utf-8").read()
check("plan.md lists the correction as read -> corrected", "Corrections you made to what David read" in plan_fix and "Sales:" in plan_fix and "(read) ->" in plan_fix and "(you)" in plan_fix)
ap_fix = json.load(open(os.path.join(DAYS, DAY, "approval.json"), encoding="utf-8"))
check("the approval records the corrections, so the audit trail is in the approval itself", ap_fix["corrections"] and "Sales" in ap_fix["corrections"][0])
exp_rows = [{"label": e["label"], "amount": e["amount"], "safe": False} for e in REAL["cash"]["expenses"]]
exp_rows[0]["amount"] = round(exp_rows[0]["amount"] + 5, 2)
exp_rows.append({"label": "new item", "amount": 7.5, "safe": True})
conf_e = {k: v for k, v in allc.items() if not k.startswith("expense.")}
for i in range(1, len(exp_rows) - 1):
    conf_e[f"expense.{i}"] = True
r_exp = D.finish(DAY, answers=dict(base, confirmed=conf_e, overrides={"expenses": exp_rows}), **kw)
apply_exp = open(os.path.join(r_exp["pack_dir"], "apply.sql"), encoding="utf-8").read()
check("an expense can be changed and a new one added 'paid from the safe': the SQL has the new amount and the new row as 'safe'; the changed row settles itself while the unchanged ones were ticked",
      r_exp["ok"] and f"{exp_rows[0]['amount']:g}" in apply_exp and "'new item', 7.5, 'safe'" in apply_exp, str(r_exp.get("reasons")))
left_e = RC.unsettled(qq, {k: v for k, v in conf_e.items() if k != "expense.1"}, {"expenses": exp_rows})
check("an unchanged expense row that is not ticked is still open, and the changed one is not", left_e == [f"Expense: {qq['expenses'][1]['label']} {qq['expenses'][1]['amount']:g}"], str(left_e))
conf_m = {k: v for k, v in allc.items() if not k.startswith("meat.beef")}
ov_m = {"meat": {"beef": {"received_kg": 2, "spit_kg": 7.7, "sold_kg": 3.0, "waste_kg": 0.5, "closing": {"trim_kg": 1.0, "main_kg": 30.0, "tallow_kg": 8.0}}}}
r_m = D.finish(DAY, answers=dict(base, confirmed=conf_m, overrides=ov_m), **kw)
apply_m = open(os.path.join(r_m["pack_dir"], "apply.sql"), encoding="utf-8").read()
check("meat corrections reach the SQL: beef received 2 (the Sep 23 tallow-only case), sold 3, waste 500 g, closing 1 / 30 / 8",
      r_m["ok"] and "'beef', 2)" in apply_m and ", 3, 0, 1000, 0)" in apply_m and "'beef', 500)" in apply_m and "'beef', 1, 30, 8)" in apply_m, str(r_m.get("reasons")))
plan_m = open(os.path.join(r_m["pack_dir"], "plan.md"), encoding="utf-8").read()
check("and the plan lists them (received, closing weights)", "beef received" in plan_m and "beef closing weights" in plan_m)
for bad_ov, why in (({"cash": {"pos_sales": -5}}, "negative"), ({"cash": {"pos_sales": "9"}}, "text"), ({"cash": {"nope": 1}}, "unknown field"), ({"expenses": [{"label": "", "amount": 1}]}, "blank name"), ({"meat": {"fish": {}}}, "unknown meat")):
    try:
        D.finish(DAY, answers=dict(base, overrides=bad_ov), **kw)
        ok = False
    except D.DayError as e:
        ok = "not valid" in str(e)
    check(f"a correction that is not valid ({why}) is refused before anything is made", ok)
held_line = qq["held_lines"][0]
conf_l = {k: v for k, v in allc.items() if k != f"line.{held_line['line']}"}
r_l = D.finish(DAY, answers=dict(base, confirmed=conf_l, overrides={"lines": {str(held_line["line"]): {"item": "C20", "qty": 7}}}), **kw)
la = json.load(open(os.path.join(DAYS, DAY, "manager_line_answers.json"), encoding="utf-8"))["answers"]
check("a held line corrected to another item and quantity is recorded for the learning step with the CORRECTED item and quantity, and its other spelling too",
      r_l["ok"] and la[0]["item"] == "C20" and la[0]["qty"] == 7 and la[0]["answer"] == "corrected" and la[1]["name"] == "سعتنتز" and la[1]["item"] == "C20", str(la)[:300])
r_l2 = D.finish(DAY, answers=dict(base, overrides={}), **kw)
la2 = json.load(open(os.path.join(DAYS, DAY, "manager_line_answers.json"), encoding="utf-8"))["answers"]
check("a held line that is only confirmed is recorded as 'ok' with David's guess", la2[0]["answer"] == "ok" and la2[0]["item"] == held_line["item"] and "learn_command" in r_l2)

r_c = D.finish(DAY, answers=dict(base, confirmed=some, overrides={"cash": {"pos_sales": REAL["cash"]["pos_sales"] + 10}}), **kw)
con2 = RH.make_db(S.load_backup(BK0))
RH.run(con2, open(os.path.join(r_c["pack_dir"], "apply.sql"), encoding="utf-8").read())
snap3 = RH.dump(con2)
BK3 = os.path.join(ROOT, "app", "backup", "2026-09-25")
empty_backup_dir(BK3, extra={t: [{k: (bool(v) if k in BOOL else v) for k, v in row.items() if v is not None} for row in snap3[t]] for t in snap3})
ck3 = D.check(DAY, **dict(kw, backup=BK3))
check("check uses the corrected copy when the approval belongs to it: the app that holds the corrected rows 'matches'", ck3["ok"] and ck3["status"] == "matches", str(ck3))

r_n = D.finish(DAY, answers=dict(base, notes={"cash.pos_sales": "  the Helper and the receipt agree  ", "check.0": "I phoned the manager of the shift", "line.1": "", "nonsense.key": "kept as is"}), **kw)
plan_n = open(os.path.join(r_n["pack_dir"], "plan.md"), encoding="utf-8").read()
ap_n = json.load(open(os.path.join(DAYS, DAY, "approval.json"), encoding="utf-8"))
check("a note written beside an item goes into the approval and the plan under its label, trimmed; an empty note is dropped", r_n["ok"] and "## Your notes" in plan_n and "- Sales: the Helper and the receipt agree" in plan_n
      and any(n.startswith("Warning:") and n.endswith("I phoned the manager of the shift") for n in ap_n["manager_notes"]) and not any("line.1" in n or n.endswith(": ") for n in ap_n["manager_notes"]) and "nonsense.key: kept as is" in ap_n["manager_notes"], str(ap_n.get("manager_notes")))
r_n0 = D.finish(DAY, answers=base, **kw)
check("with no notes the plan has no notes section and the approval no manager_notes", "Your notes" not in open(os.path.join(r_n0["pack_dir"], "plan.md"), encoding="utf-8").read() and "manager_notes" not in json.load(open(os.path.join(DAYS, DAY, "approval.json"), encoding="utf-8")))

# ── the command line ─────────────────────────────────────────────────────────────────────────────────────────────────────────────
check("a bad date on the command line exits 2 with no writes", D.main(["prepare", "2026-13-40"]) == 2)
code = "\n".join(l for l in inspect.getsource(D).splitlines() if not l.lstrip().startswith("#"))
check("the runner never connects to anything: no http, socket, supabase, psycopg or sqlite3 import in it", not re.search(r"(?m)^\s*(import|from)\s+(requests|httpx|socket|urllib|sqlite3|psycopg|supabase)", code))
check("it reads the app's env file for the project address only (a regex for https://<ref>.supabase.co), never a key", "supabase.co" in code and not re.search(r"(?i)service_role|SECRET_KEY\s*=|ANON_KEY|api[_-]?key", code.split("steps_markdown")[0]))

print("\n-- an expense whose amount the photo reading could not find (seen live on 2026-10-12) --")
pkg_none = {"business_day": "2026-10-12", "flags": [], "extraction": {"summary_by_items": {"lines": []}}, "computed": {"expected": 100.0, "counted": 100.0, "diff": 0.0}, "cfo_results": {"meat": {}},
            "cash": {"old_cash": 10.0, "pos_sales": 20.0, "wallet": 0.0, "transfer_to_admin": 5.0, "expenses": [{"label": "veg.", "amount": 282.22}, {"label": "invoices", "amount": None}], "actual_cash_count": 1.0}}
try:
    qn = D.questions_for(pkg_none, [], [], {"days": [], "cash_sheets": []}, None, None)
    ok_q = True
except Exception as ex:
    qn, ok_q = None, False
check("a missing expense amount no longer crashes the question builder", ok_q)
check("it is left out of the total, shown as 'no amount read', and the manager is asked for it", ok_q and qn["figures"]["expenses_total"] == 282.22 and qn["expenses"][1]["amount"] is None and not qn["expenses"][1]["big"]
      and any("invoices" in w and "no amount" in w for w in qn["warnings"]), str(qn and qn["warnings"]))
check("the markdown for the manager says (no amount read), never None", ok_q and "(no amount read)" in D.questions_markdown("2026-10-12", qn, "x.json") and "invoices None" not in D.questions_markdown("2026-10-12", qn, "x.json"))

qq_none = {"figures": dict({k: 0 for k in ("old_cash", "sales", "wallet", "expenses_total", "transfer", "expected", "counted", "diff")}, chicken="c", beef="b"), "checks": [], "blockers": [], "held_lines": [{"line": 7, "name": "x", "alt": None, "qty": None, "qty_alt": None, "guess": "y", "item": "C01", "score": 90, "status": "suggested"},
                                                                    {"line": 8, "name": "x2", "alt": "x3", "qty": 2.0, "qty_alt": None, "guess": "y", "item": "C01", "score": 90, "status": "suggested"}],
           "deliveries": {"chicken": 1.0, "beef": 1.0}, "expenses": [], "notes": [], "warnings": []}
try:
    md_none = D.questions_markdown("2026-10-15", qq_none, "DRAFT-x.json")
except Exception as e:
    md_none = repr(e)
check("an item line whose quantity the photo reading could not give does not stop the day (it shows x?), as on 2026-10-15", "line 7: 'x' x?" in md_none and "line 8: 'x2' / 'x3' x2" in md_none)


shutil.rmtree(TMP, ignore_errors=True)
print(f"\n  {sum(results)}/{len(results)} passed")
sys.exit(0 if all(results) else 1)
