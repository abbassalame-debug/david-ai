"""
test_restaurant_upload.py -- offline tests for restaurant_upload.py (the manager's "Upload to the app" button) and its API endpoints.

The "database" is a SQLite stand-in with the same tables, driven through the same tiny connection interface the real Postgres connection uses; the Demo Bistro tree, the
backup and the folders are temp folders; the photo reading is replaced by the saved demo draft (examples/demo_bistro). Nothing here touches Supabase, a network or a real folder.
Run: py -3.11 test_restaurant_upload.py
"""
from david_config import DAVID_DATA, DAVID_HOME
import sys, os, json, glob, re, copy, shutil, tempfile, inspect, warnings
sys.stdout.reconfigure(encoding="utf-8")
sys.path.insert(0, str(DAVID_HOME))
sys.path.insert(0, str(DAVID_HOME / "ui"))
warnings.filterwarnings("ignore")

PASS = "\033[92m  PASS\033[0m"
FAIL = "\033[91m  FAIL\033[0m"
results = []


def check(name, ok, detail=""):
    results.append(bool(ok))
    print(f"{PASS if ok else FAIL}  {name}" + (f" -- {detail}" if detail else ""))


import openpyxl
from fastapi import FastAPI
from starlette.testclient import TestClient
import restaurant_corrections as RC
import restaurant_day as D
import restaurant_pipeline as P
import restaurant_rehearsal as RH
import restaurant_sql as S
import restaurant_upload as UP
import restaurant_api as A

TMP = tempfile.mkdtemp(prefix="upload_test_")
DAY = "2026-09-22"
SECRET = "Sup3rS3cretPw!"
BOOL = ("closed", "float_open_ok", "float_close_ok")

# ── the statement splitter and the file shape ──────────────────────────────────────────────────────────────────────────────────────
sql = """-- header; with a semicolon
begin;
insert into t (a) values ('x; y'), ('it''s; fine');   -- trailing; comment
update t set a = 'q' where id = 1;
commit;

select 'a' as what, 1 as rows_found, 1 as rows_expected
union all
select 'b', 1, 1;
"""
st = UP.statements(sql)
check("the splitter cuts on ';' only outside quotes, keeps '' escapes, and drops comments", st == ["begin", "insert into t (a) values ('x; y'), ('it''s; fine')", "update t set a = 'q' where id = 1", "commit", "select 'a' as what, 1 as rows_found, 1 as rows_expected\nunion all\nselect 'b', 1, 1"], str(st))
body, chk = UP.split_file(sql)
check("a generated file splits into its body (no begin / commit) and its check table", body == st[1:3] and chk.startswith("select 'a' as what"))
for bad_sql in ("begin;\ncommit;\n", "begin;\ninsert into t values (1);\ncommit;\n", "select 1;"):
    try:
        UP.split_file(bad_sql)
        ok = False
    except UP.UploadError:
        ok = True
    check(f"a file without statements or without its check table is refused: {bad_sql[:24]!r}", ok)

# ── secrets ──────────────────────────────────────────────────────────────────────────────────────────────────────────────────────
url = f"postgresql://postgres.abc:{SECRET}@aws-0.pooler.supabase.com:5432/postgres"
old = os.environ.pop(UP.ENV_URL, None)
check("without the variable the upload is off", UP.configured() is False)
try:
    UP.default_connect()
    ok = False
except UP.UploadError as e:
    ok = UP.ENV_URL in str(e)
check("the error names the variable to set", ok)
os.environ[UP.ENV_URL] = url
check("with the variable set the button is on", UP.configured() is True)
e1 = UP.clean_error(RuntimeError(f"could not connect to {url} password {SECRET} failed"))
check("an error never carries the connection string or the password (not as typed, not inside another URL)", SECRET not in e1 and url not in e1 and "***" in e1, e1)
e2 = UP.clean_error(RuntimeError("bad postgres://user:pa%24%24@host/db thing"), url)
check("any postgres:// address in an error is blanked", "pa%24%24" not in e2 and "postgres://***" in e2)
os.environ.pop(UP.ENV_URL)
if old is not None:
    os.environ[UP.ENV_URL] = old

# ── a stand-in database behind the same interface ──────────────────────────────────────────────────────────────────────────────────
class Fake:
    def __init__(self, backup, fail_on=None, leak=None):
        self.con = RH.make_db(backup)
        self.con.isolation_level = None
        self.tx = False
        self.fail_on, self.leak = fail_on, leak
        self.closed = False

    def _begin(self):
        if not self.tx:
            self.con.execute("begin")
            self.tx = True

    def exec(self, s):
        if s.lower().startswith("set transaction"):
            return
        self._begin()
        if self.fail_on and self.fail_on in s:
            raise RuntimeError(f"insert failed on connection {self.leak or ''}")
        self.con.execute(s)

    def query(self, s):
        self._begin()
        cur = self.con.execute(s)
        names = [d[0] for d in cur.description]
        return names, [tuple(bool(v) if n in BOOL and v is not None else v for n, v in zip(names, r)) for r in cur.fetchall()]

    def commit(self):
        if self.tx:
            self.con.execute("commit")
            self.tx = False

    def rollback(self):
        if self.tx:
            self.con.execute("rollback")
            self.tx = False

    def close(self):
        self.closed = True


# ── the day's files, made the way the runner makes them ────────────────────────────────────────────────────────────────────────────
ROOT = os.path.join(TMP, "abo")
FOLDER = os.path.join(ROOT, "09-2026", "22-09-2026")
os.makedirs(FOLDER)
os.makedirs(os.path.join(ROOT, "app"))
wb = openpyxl.Workbook()
wb.remove(wb.active)
ws = wb.create_sheet("22-09-2026")
ws.cell(3, 3).value = "قبل الشك"
for c, v in ((4, 4.0), (5, 150.0), (15, 1.0), (16, 6.0), (17, 28.0)):
    ws.cell(3, c).value = v
ws.cell(6, 2).value = "الوزن السيخ"; ws.cell(6, 4).value = 35.68
ws.cell(6, 14).value = "الوزن السيخ"; ws.cell(6, 17).value = 7.7
ws.cell(20, 2).value = "مجموع السياخ دجاج"; ws.cell(22, 2).value = 35.68
ws.cell(20, 14).value = "مجموع السياخ لحم"; ws.cell(22, 14).value = 7.7
ws.cell(30, 1).value = "end-of-shift closing summary"
ws.cell(31, 3).value = "قبل الشك"
for c, v in ((4, 5.0), (5, 182.0), (15, 1.5), (16, 6.0), (17, 28.5)):
    ws.cell(31, c).value = v
wb.save(os.path.join(FOLDER, "فوارق الشك مع الاوزان 22-09-2026.xlsx"))
REAL = json.load(open(str(DAVID_HOME / "examples" / "demo_bistro" / "draft_package.demo.json"), encoding="utf-8"))   # the demo draft package (synthetic)


def fake_run_day(folder, *, business_day=None, drafts_dir=None, **kw):
    pkg = copy.deepcopy(REAL)
    pkg["draft_id"] = f"DRAFT-{business_day}-{len(glob.glob(os.path.join(drafts_dir, f'DRAFT-{business_day}-*.json'))) + 1:03d}"
    path = os.path.join(drafts_dir, pkg["draft_id"] + ".json")
    json.dump(pkg, open(path, "w", encoding="utf-8"), ensure_ascii=False)
    return {"ok": True, "package": pkg, "package_path": path, "validation": {"valid": True, "errors": []}, "review_path": None}


P.run_day = fake_run_day
BK = os.path.join(ROOT, "app", "backup", "2026-09-22")
os.makedirs(BK)
for t in S.TABLES:
    json.dump([], open(os.path.join(BK, t + ".json"), "w", encoding="utf-8"))
DAYS, REV = os.path.join(TMP, "days"), os.path.join(TMP, "review")
kw = dict(restaurant_root=ROOT, backup=BK, days_dir=DAYS, review_dir=REV)
D.prepare(DAY, log=lambda m: None, **kw)
qq = D._questions(DAY, DAYS)
answers = {"approved_by": "Demo User", "confirmed": {k: True for k, _ in RC.required_items(qq)}, "cleared_blockers": {str(i): "looked" for i in range(1, len(qq["blockers"]) + 1)}, "paid_from_safe": [], "allow_updates": [], "overrides": {}}
fin = D.finish(DAY, answers=answers, **kw)
APPLY = open(os.path.join(REV, f"day_{DAY}", "apply.sql"), encoding="utf-8").read()
SHA = UP.sha(APPLY)
EMPTY = S.load_backup(BK)
check("the day's SQL exists, made by the runner", fin["ok"] and fin["status"] == "sql_ready")


def fresh(**k):
    return Fake(EMPTY, **k)


def up(conn, dry=False, **over):
    args = dict(dry_run=dry, sql_sha=SHA, confirm_day=DAY, days_dir=DAYS, review_dir=REV, connect=lambda: conn)
    args.update(over)
    return UP.upload(DAY, **args)


def refused(fn):
    try:
        fn()
        return None
    except UP.UploadError as e:
        return str(e)


# ── a dry run runs on the database and saves nothing ─────────────────────────────────────────────────────────────────────────────
db = fresh()
before = RH.dump(db.con)
r = up(db, dry=True)
check("a dry run runs the real file, finds every check line, and reports it", r["ok"] and r["status"] == "dry_run_ok" and r["committed"] is False and r["dry_run"] is True and r["checked"] >= 12 and r["not_found"] == [], str(r))
check("and NOTHING is saved: every table is exactly as before and the connection was closed", RH.dump(db.con) == before and db.closed)
a0 = json.load(open(os.path.join(DAYS, DAY, "upload.json"), encoding="utf-8"))
check("the dry run is recorded as a dry run (committed false) so it can never be mistaken for an upload", a0["dry_run"] is True and a0["committed"] is False and a0["sql_sha"] == SHA)

# ── the real upload ──────────────────────────────────────────────────────────────────────────────────────────────────────────────
db = fresh()
r = up(db)
rows = RH.dump(db.con)
check("a real upload commits, and the day's rows are in the database exactly as the file says", r["ok"] and r["status"] == "committed" and r["committed"] is True and len(rows["cash_sheets"]) == 1 and len(rows["days"]) == 1
      and rows["cash_sheets"][0]["sales_usd"] == REAL["cash"]["pos_sales"] and len(rows["spit_loads"]) == 2 and len(rows["meat_closing_weights"]) == 2, str(r))
check("it then re-reads the live tables and re-plans: the day is 'nothing to do', which is the proof the app holds exactly the draft", r["verified"] is True and r["differences"] == [])
a1 = json.load(open(os.path.join(DAYS, DAY, "upload.json"), encoding="utf-8"))
check("the audit record names the SQL (its fingerprint), who approved it, when, and that it was committed", a1["committed"] is True and a1["sql_sha"] == SHA and a1["approved_by"] == "Demo User" and a1["verified"] is True and re.match(r"\d{4}-\d{2}-\d{2}T", a1["at"]))
check("the same SQL cannot be uploaded twice (the second press is refused, naming the time)", "already uploaded" in (refused(lambda: up(db)) or ""))
r2 = up(fresh(), dry=True)
check("a dry run after a real upload is still allowed (it does not overwrite the record of the real one)", r2["ok"] and json.load(open(os.path.join(DAYS, DAY, "upload.json"), encoding="utf-8"))["committed"] is True)
chk_live = UP.live_check(DAY, days_dir=DAYS, connect=lambda: db)
check("the live check says the app matches the draft, with no backup export needed", chk_live["ok"] and chk_live["status"] == "matches" and chk_live["source"] == "the live database")
import menu_tools as _MT
_real_validate = _MT.validate_draft_package
_MT.validate_draft_package = lambda **kw: {"valid": False, "errors": [{"code": "MATCH_NOT_BEST", "message": "another item scores higher: a spelling learned after the approval"}]}
try:
    chk_grown = UP.live_check(DAY, days_dir=DAYS, connect=lambda: db)
finally:
    _MT.validate_draft_package = _real_validate
check("re-checking an APPROVED day does not re-validate its draft against the menu memory as it is now (learning a spelling after the approval used to make the check refuse the day)", chk_grown["ok"] and chk_grown["status"] == "matches")
chk_before = UP.live_check(DAY, days_dir=DAYS, connect=lambda: fresh())
check("the live check on a database without the day lists what is missing", chk_before["ok"] is False and chk_before["status"] == "sql_ready" and len(chk_before["differences"]) >= 12)
os.remove(os.path.join(DAYS, DAY, "upload.json"))

# ── the gates ────────────────────────────────────────────────────────────────────────────────────────────────────────────────────
check("a wrong confirmation day is refused before the database is touched", "not this day" in (refused(lambda: up(fresh(), confirm_day="2026-09-23")) or ""))
db_g = fresh()
msg = refused(lambda: up(db_g, sql_sha="0" * 64))
check("SQL that is not the SQL on disk (the page shows an old version) is refused", "not the SQL on disk" in (msg or "") and RH.dump(db_g.con) == RH.dump(fresh().con))
check("no apply.sql for the day is refused", "make the SQL files first" in (refused(lambda: UP.upload("2026-09-30", dry_run=True, sql_sha=SHA, confirm_day="2026-09-30", days_dir=DAYS, review_dir=REV, connect=lambda: fresh())) or ""))
try:
    UP.upload("nope", dry_run=True, sql_sha=SHA, confirm_day="nope", days_dir=DAYS, review_dir=REV)
    ok = False
except D.DayError:
    ok = True
check("a day that is not a date is refused", ok)

# the app changed since the SQL was made
manual = copy.deepcopy(EMPTY)
sheet = {"id": "bbbbbbbb-0000-4000-8000-000000000001", "day_id": S.stable_id(f"day:{S.BRANCH_ID}:{DAY}"), "sales_usd": 1.0, "count_cash_usd": 0, "count_cash_lbp": 0, "count_wallet_usd": 0, "rate_used": 90000,
         "safe_old_usd": 0, "safe_counted_usd": 0, "next_float_usd": 211.11, "float_open_ok": False, "float_close_ok": False, "transfer_to_admin_usd": 0, "closed": True}
manual["days"].append({"id": sheet["day_id"], "branch_id": S.BRANCH_ID, "on_date": DAY, "status": "ended", "rate_usd_to_lbp": 90000})
manual["cash_sheets"].append(sheet)
db_m = Fake(manual)
before_m = RH.dump(db_m.con)
msg = refused(lambda: up(db_m))
check("if someone entered the day by hand after the SQL was made, the upload stops and says the app changed (nothing is written)", msg is not None and "changed" in msg and RH.dump(db_m.con) == before_m, str(msg))
db_same = Fake(S.load_backup(BK))
RH.run(db_same.con, APPLY)

r_same = up(db_same)
check("if the app already holds exactly this day there is nothing to upload (and it says so)", r_same["status"] == "already_in_app" and r_same["committed"] is False)

# a database error rolls everything back and never shows the secret
os.environ[UP.ENV_URL] = url
db_f = fresh(fail_on="insert into meat_receipts", leak=url)
before_f = RH.dump(db_f.con)
msg = refused(lambda: up(db_f))
check("a failing statement halfway rolls EVERYTHING back (the earlier inserts are gone too), and the error does not contain the connection string or password",
      msg is not None and "nothing was saved" in msg and SECRET not in msg and url not in msg and RH.dump(db_f.con) == before_f, str(msg))
os.environ.pop(UP.ENV_URL)
check("the connection is closed after a failure", db_f.closed)
UP._LOCK.acquire()
check("two uploads at once are refused (one runs at a time)", "another upload is running" in (refused(lambda: up(fresh())) or ""))
UP._LOCK.release()
got = UP._LOCK.acquire(blocking=False)
if got:
    UP._LOCK.release()
check("the lock is free again after a refusal", got)

# the connection test
tc = UP.test_connection(lambda: fresh())
check("the connection test reads only: it reports how many days the app holds and the newest", tc["ok"] and tc["days"] == 0 and tc["last_day"] is None)
tc2 = UP.test_connection(lambda: db)
check("it reports a database that has days", tc2["ok"] and tc2["days"] == 1 and tc2["last_day"] == DAY)
def boom():
    raise UP.UploadError("could not connect: timeout")
check("a connection that fails gives a message, not a crash", UP.test_connection(boom) == {"ok": False, "message": "could not connect: timeout"})

# ── the API ──────────────────────────────────────────────────────────────────────────────────────────────────────────────────────
conn_box = {"db": None}
app = FastAPI()
app.include_router(A.make_router({"days_dir": DAYS, "review_dir": REV, "inbox_dir": os.path.join(TMP, "inbox"), "restaurant_root": ROOT, "backup": BK, "connect": lambda: conn_box["db"]}))
c = TestClient(app, base_url="http://127.0.0.1:8000")
conn_box["db"] = fresh()
check("the status says the upload is switched on (a connection was provided) and which variable switches it on", c.get("/api/restaurant/db/status").json()["configured"] is True and c.get("/api/restaurant/db/status").json()["variable"] == UP.ENV_URL)
check("the connection test works through the API", c.post("/api/restaurant/db/test").json()["ok"] is True)
d = c.get(f"/api/restaurant/day/{DAY}").json()
check("the day view carries the fingerprint of the SQL on disk (the page sends it back with the upload) and no upload yet", d["apply_sha"] == SHA and d["upload"] is None)
check("an upload request without its fields, or with the wrong kinds, is refused (400)", c.post(f"/api/restaurant/day/{DAY}/upload", json={}).status_code == 400 and c.post(f"/api/restaurant/day/{DAY}/upload", json={"dry_run": "yes", "confirm_day": DAY, "sql_sha": SHA}).status_code == 400)
r = c.post(f"/api/restaurant/day/{DAY}/upload", json={"dry_run": False, "confirm_day": "2026-09-23", "sql_sha": SHA})
check("a refusal comes back as 422 with the message the page shows", r.status_code == 422 and r.json()["status"] == "refused" and "not this day" in r.json()["message"])
conn_box["db"] = fresh()
r = c.post(f"/api/restaurant/day/{DAY}/upload", json={"dry_run": True, "confirm_day": DAY, "sql_sha": SHA})
check("a dry run through the API", r.status_code == 200 and r.json()["status"] == "dry_run_ok")
conn_box["db"] = fresh()
r = c.post(f"/api/restaurant/day/{DAY}/upload", json={"dry_run": False, "confirm_day": DAY, "sql_sha": SHA})
check("a real upload through the API commits and verifies", r.status_code == 200 and r.json()["status"] == "committed" and r.json()["verified"] is True)
d2 = c.get(f"/api/restaurant/day/{DAY}").json()
check("the day view then shows the upload record", d2["upload"]["committed"] is True and d2["upload"]["sql_sha"] == SHA)
ck = c.post(f"/api/restaurant/day/{DAY}/check?live=true").json()
check("the check can read the live database instead of a backup", ck["ok"] is True and ck["source"] == "the live database")
import server as srv
hud = TestClient(srv.app, base_url="http://127.0.0.1:8000")
check("a page on another site cannot press the upload (the HUD's origin guard says 403)", hud.post(f"/api/restaurant/day/{DAY}/upload", json={"dry_run": False, "confirm_day": DAY, "sql_sha": SHA}, headers={"Origin": "http://evil.example"}).status_code == 403
      and hud.post("/api/restaurant/db/test", headers={"Origin": "http://evil.example"}).status_code == 403)
check("the real HUD says the upload is off unless the connection variable is set", hud.get("/api/restaurant/db/status").json()["configured"] is bool(os.environ.get(UP.ENV_URL, "").strip()))

code = "\n".join(l for l in inspect.getsource(UP).splitlines() if not l.lstrip().startswith("#"))
check("it never prints or logs anything (so the connection string cannot be shown)", "print(" not in code and "logging" not in code and "log(" not in code)
check("the connection string is read in exactly one place (the environment variable), never from a file, and only used to connect", code.count("os.environ.get(ENV_URL") == 3 and "open(" in code and not re.search(r"open\([^)]*(env|\.env)", code))
after_doc = code.split('"""')[2]
check("it never runs undo.sql, never deletes, never drops: the only SQL it runs is the statements of the approved apply.sql and its own reads", "undo" not in after_doc and not re.search(r"(?i)\b(drop|delete|truncate)\s", after_doc))

shutil.rmtree(TMP, ignore_errors=True)
print(f"\n  {sum(results)}/{len(results)} passed")
sys.exit(0 if all(results) else 1)
